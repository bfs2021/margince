"""Margince-branded PowerPoint slides on python-pptx.

Every content slide carries the emerald band with the white lockup at the top.
Import the helpers, or run the file to write a sample deck:

    python3 margince_pptx.py sample.pptx
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ASSETS = Path(__file__).resolve().parent.parent / "assets"

EMERALD = RGBColor(0x0B, 0x7A, 0x53)
EMERALD_TEXT = RGBColor(0x0A, 0x6F, 0x4B)
EMERALD_BG = RGBColor(0xE8, 0xF3, 0xEE)
PAPER = RGBColor(0xF1, 0xF5, 0xF2)
SURFACE = RGBColor(0xFB, 0xFC, 0xFB)
INK = RGBColor(0x10, 0x1A, 0x15)
INK2 = RGBColor(0x33, 0x40, 0x3A)
INK3 = RGBColor(0x66, 0x73, 0x6C)
LINE = RGBColor(0xD9, 0xDF, 0xDB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# PowerPoint swaps in a system face when these are not installed on the viewing machine.
HEADING_FONT = "Outfit"
BODY_FONT = "Geist"
LABEL_FONT = "Geist Mono"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
BAND_H = Inches(0.8)
MARGIN = Inches(0.75)


def new_deck():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _blank(prs, ground):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = ground
    return slide


def _flat(shape):
    # The default theme style adds a drop shadow that LibreOffice draws even
    # when the shape's own effect list is empty, so the style reference goes.
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    return shape


def _rect(slide, left, top, width, height, color):
    shape = _flat(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def _text(slide, left, top, width, height, text, *, font=BODY_FONT, size=20,
          color=INK2, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          tracking=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text.upper() if tracking else text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if tracking:
        run.font._element.set("spc", "100")
    return box


def _lockup(slide, left, top, height, white=True):
    name = "lockup-white.png" if white else "lockup-color.png"
    return slide.shapes.add_picture(str(ASSETS / name), left, top, height=height)


def _band(slide, label):
    _rect(slide, 0, 0, SLIDE_W, BAND_H, EMERALD)
    _lockup(slide, MARGIN, Inches(0.24), Inches(0.32))
    if label:
        _text(slide, Inches(7.5), 0, SLIDE_W - Inches(7.5) - MARGIN, BAND_H, label,
              font=LABEL_FONT, size=11, color=WHITE, align=PP_ALIGN.RIGHT,
              anchor=MSO_ANCHOR.MIDDLE, tracking=True)


def _footer(slide, number):
    top = SLIDE_H - Inches(0.55)
    _rect(slide, MARGIN, top, SLIDE_W - 2 * MARGIN, Pt(0.75), LINE)
    _text(slide, MARGIN, top + Inches(0.1), Inches(4), Inches(0.3), "Margince", size=11, color=INK3)
    _text(slide, SLIDE_W - MARGIN - Inches(1), top + Inches(0.1), Inches(1), Inches(0.3),
          str(number), size=11, color=INK3, align=PP_ALIGN.RIGHT)


def title_slide(prs, title, subtitle=None, byline=None):
    slide = _blank(prs, EMERALD)
    _lockup(slide, Inches(0.9), Inches(0.7), Inches(0.5))
    _text(slide, Inches(0.9), Inches(3.6), Inches(11), Inches(1.6), title,
          font=HEADING_FONT, size=56, color=WHITE, bold=True, anchor=MSO_ANCHOR.BOTTOM)
    if subtitle:
        _text(slide, Inches(0.9), Inches(5.3), Inches(10), Inches(0.6), subtitle, size=22, color=WHITE)
    if byline:
        _text(slide, Inches(0.9), Inches(6.4), Inches(10), Inches(0.4), byline,
              font=LABEL_FONT, size=12, color=EMERALD_BG, tracking=True)
    return slide


def section_slide(prs, number, title):
    slide = _blank(prs, PAPER)
    strip_w = Emu(int(SLIDE_W * 0.34))
    _rect(slide, 0, 0, strip_w, SLIDE_H, EMERALD)
    slide.shapes.add_picture(str(ASSETS / "logomark-white.png"), Inches(0.6), Inches(0.6), height=Inches(0.6))
    left = strip_w + Inches(0.9)
    _text(slide, left, Inches(2.4), Inches(7), Inches(1.3), f"{number:02d}",
          font=HEADING_FONT, size=88, color=EMERALD, bold=True, anchor=MSO_ANCHOR.BOTTOM)
    _text(slide, left, Inches(3.8), Inches(7.5), Inches(1.4), title,
          font=HEADING_FONT, size=44, color=INK, bold=True)
    return slide


def content_slide(prs, title, *, bullets=None, eyebrow=None, lead=None, label=None):
    """A banded slide; the caller may add shapes under the head, from body_top() down."""
    slide = _blank(prs, PAPER)
    _band(slide, label)
    top = BAND_H + Inches(0.5)
    width = SLIDE_W - 2 * MARGIN
    if eyebrow:
        _text(slide, MARGIN, top, width, Inches(0.3), eyebrow,
              font=LABEL_FONT, size=12, color=EMERALD_TEXT, tracking=True)
        top += Inches(0.4)
    _text(slide, MARGIN, top, width, Inches(0.8), title, font=HEADING_FONT, size=36, color=INK, bold=True)
    top += Inches(0.85)
    if lead:
        _text(slide, MARGIN, top, width, Inches(0.5), lead, size=20, color=INK3)
        top += Inches(0.65)
    if bullets:
        _bullets(slide, MARGIN, top + Inches(0.15), width, bullets)
    _footer(slide, len(prs.slides))
    return slide


def body_top():
    return BAND_H + Inches(2.4)


def _bullets(slide, left, top, width, bullets):
    for i, point in enumerate(bullets):
        y = top + Inches(0.55) * i
        _rect(slide, left, y + Inches(0.1), Inches(0.13), Inches(0.13), EMERALD)
        _text(slide, left + Inches(0.35), y, width - Inches(0.35), Inches(0.5), point, size=20)


def card(slide, left, top, width, height, heading, body):
    box = _flat(slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height))
    box.adjustments[0] = 0.08
    box.fill.solid()
    box.fill.fore_color.rgb = SURFACE
    box.line.color.rgb = LINE
    pad = Inches(0.3)
    _text(slide, left + pad, top + pad, width - 2 * pad, Inches(0.5), heading,
          font=HEADING_FONT, size=22, color=INK, bold=True)
    _text(slide, left + pad, top + pad + Inches(0.55), width - 2 * pad, height - Inches(1.1), body, size=17)


def stat(slide, left, top, width, value, label):
    _rect(slide, left, top, width, Pt(3), EMERALD)
    _text(slide, left, top + Inches(0.2), width, Inches(0.9), value,
          font=HEADING_FONT, size=54, color=INK, bold=True)
    _text(slide, left, top + Inches(1.15), width, Inches(0.5), label, size=16, color=INK3)


def closing_slide(prs, headline, line=None):
    slide = _blank(prs, EMERALD)
    _lockup(slide, Inches(0.9), Inches(0.7), Inches(0.5))
    _text(slide, Inches(0.9), Inches(4.2), Inches(11), Inches(1.2), headline,
          font=HEADING_FONT, size=48, color=WHITE, bold=True, anchor=MSO_ANCHOR.BOTTOM)
    if line:
        _text(slide, Inches(0.9), Inches(5.5), Inches(10), Inches(0.5), line, size=20, color=WHITE)
    return slide


def _sample(path):
    prs = new_deck()
    title_slide(prs, "Deck title in one line",
                "One sentence that says what the reader gets from this deck.",
                "Author name · 23 October 2026")
    section_slide(prs, 1, "Section title")
    content_slide(prs, "A heading that states one idea.", eyebrow="Eyebrow label",
                  lead="A lead sentence that frames the points below.",
                  bullets=["First point, one sentence of detail.",
                           "Second point, one sentence of detail.",
                           "Third point, one sentence of detail."],
                  label="Deck title")
    slide = content_slide(prs, "Two options, side by side.", eyebrow="Comparison", label="Deck title")
    half = (SLIDE_W - 2 * MARGIN - Inches(0.4)) // 2
    card(slide, MARGIN, Inches(2.6), half, Inches(2), "Option A", "What it is, what it costs, what it gives.")
    card(slide, MARGIN + half + Inches(0.4), Inches(2.6), half, Inches(2), "Option B",
         "What it is, what it costs, what it gives.")
    slide = content_slide(prs, "Three numbers the reader should keep.", eyebrow="In figures", label="Deck title")
    third = (SLIDE_W - 2 * MARGIN - Inches(0.8)) // 3
    for i, (value, label) in enumerate([("42", "What this figure counts"), ("3.5×", "What this figure counts"),
                                        ("98%", "What this figure counts")]):
        stat(slide, MARGIN + (third + Inches(0.4)) * i, Inches(2.8), third, value, label)
    closing_slide(prs, "Request a conversation", "Contact line or next step")
    prs.save(path)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: margince_pptx.py OUT.pptx")
    _sample(sys.argv[1])
