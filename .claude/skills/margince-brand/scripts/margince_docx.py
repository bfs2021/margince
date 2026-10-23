"""Margince-branded Word documents on python-docx.

Page one opens under an emerald band with the white lockup, the document kind,
the title and a meta line; every later page carries a thin emerald band. Import
the helpers, or run the file to write a sample:

    python3 margince_docx.py sample.docx
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ASSETS = Path(__file__).resolve().parent.parent / "assets"

EMERALD = "0B7A53"
EMERALD_TEXT = RGBColor(0x0A, 0x6F, 0x4B)
EMERALD_BG = "E8F3EE"
WARN_BG = "FFF6E6"
WARN_EDGE = "F5A524"
WARN_TEXT = RGBColor(0x8A, 0x53, 0x00)
INK = RGBColor(0x10, 0x1A, 0x15)
INK2 = RGBColor(0x33, 0x40, 0x3A)
INK3 = RGBColor(0x66, 0x73, 0x6C)
LINE = "D9DFDB"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Word swaps in a system face when these are not installed on the viewing machine.
HEADING_FONT = "Outfit"
BODY_FONT = "Geist"

PAGE_W = Mm(210)
GUTTER = Mm(20)


def _shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def _cell_margins(cell, top=0, bottom=0, left=0, right=0):
    props = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(int(value.twips if hasattr(value, "twips") else value)))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    props.append(margins)


def _borders(table, edges):
    """edges maps a side (top, left, bottom, right, insideH, insideV) to (hex, eighths of a point) or None."""
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{side}")
        spec = edges.get(side)
        if spec:
            node.set(qn("w:val"), "single")
            node.set(qn("w:color"), spec[0])
            node.set(qn("w:sz"), str(spec[1]))
        else:
            node.set(qn("w:val"), "nil")
        borders.append(node)
    props.append(borders)


def _bleed(table, height=None):
    """Stretch a header table from paper edge to paper edge, over the side margins."""
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    props = table._tbl.tblPr
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), str(-GUTTER.twips))
    indent.set(qn("w:type"), "dxa")
    props.append(indent)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    props.append(layout)
    table.autofit = False
    table.columns[0].width = PAGE_W
    table.rows[0].cells[0].width = PAGE_W
    if height is not None:
        table.rows[0].height = height
    _borders(table, {})


def _run(paragraph, text, *, font=BODY_FONT, size=10.5, color=INK2, bold=False, caps=False):
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.all_caps = caps
    return run


def _spacing(paragraph, before=0, after=0):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)


def _page_number(paragraph):
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        run = paragraph.add_run()
        run.font.size = Pt(8)
        run.font.color.rgb = INK3
        if kind:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        else:
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = text
        run._r.append(node)


def new_document(title, *, kind=None, meta=None):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = PAGE_W, Mm(297)
    section.left_margin = section.right_margin = GUTTER
    section.top_margin, section.bottom_margin = Mm(18), Mm(18)
    section.header_distance, section.footer_distance = Mm(0), Mm(8)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK2
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    _first_page_band(section.first_page_header, title, kind, meta)
    _thin_band(section.header)
    for footer in (section.first_page_footer, section.footer):
        para = footer.paragraphs[0]
        _run(para, "Margince    ", size=8, color=INK3)
        _page_number(para)
    return doc


def _first_page_band(header, title, kind, meta):
    table = header.add_table(rows=1, cols=1, width=PAGE_W)
    _bleed(table)
    cell = table.rows[0].cells[0]
    _shade(cell, EMERALD)
    _cell_margins(cell, top=Mm(12), bottom=Mm(10), left=GUTTER, right=GUTTER)
    logo = cell.paragraphs[0]
    logo.add_run().add_picture(str(ASSETS / "lockup-white.png"), height=Mm(6))
    _spacing(logo, after=24)
    if kind:
        para = cell.add_paragraph()
        _run(para, kind, size=8, color=WHITE, caps=True)
        _spacing(para, after=4)
    para = cell.add_paragraph()
    _run(para, title, font=HEADING_FONT, size=26, color=WHITE, bold=True)
    _spacing(para, after=6)
    if meta:
        para = cell.add_paragraph()
        _run(para, meta, size=9.5, color=WHITE)
    _above_closing_paragraph(header, table, Pt(16))


def _thin_band(header):
    table = header.add_table(rows=1, cols=1, width=PAGE_W)
    _bleed(table, height=Mm(4))
    _shade(table.rows[0].cells[0], EMERALD)
    _above_closing_paragraph(header, table, Pt(0))


def _above_closing_paragraph(header, table, gap):
    # Word rejects a header whose last block is a table, so the header's own
    # empty paragraph stays last and only sets the gap above the body.
    closing = header.paragraphs[0]
    closing._p.addprevious(table._tbl)
    closing.paragraph_format.space_after = gap
    closing.paragraph_format.line_spacing = Pt(1)


def lead(doc, text):
    para = doc.add_paragraph()
    _run(para, text, size=12.5, color=INK3)
    _spacing(para, after=12)
    return para


def heading(doc, text, level=2):
    para = doc.add_paragraph()
    if level == 2:
        props = para._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        top = OxmlElement("w:top")
        for key, value in (("val", "single"), ("sz", "12"), ("space", "6"), ("color", EMERALD)):
            top.set(qn(f"w:{key}"), value)
        border.append(top)
        props.append(border)
        _run(para, text, font=HEADING_FONT, size=16, color=INK, bold=True)
        _spacing(para, before=18, after=6)
    else:
        _run(para, text, font=HEADING_FONT, size=12, color=INK, bold=True)
        _spacing(para, before=10, after=4)
    para.paragraph_format.keep_with_next = True
    return para


def paragraph(doc, text):
    para = doc.add_paragraph()
    _run(para, text)
    return para


def bullets(doc, points):
    for point in points:
        para = doc.add_paragraph(style="List Bullet")
        _run(para, point)
        _spacing(para, after=3)


def callout(doc, label, text, *, warn=False):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    _shade(cell, WARN_BG if warn else EMERALD_BG)
    _borders(table, {"left": (WARN_EDGE if warn else EMERALD, 24)})
    _cell_margins(cell, top=Mm(3), bottom=Mm(3), left=Mm(4), right=Mm(4))
    para = cell.paragraphs[0]
    _run(para, f"{label} ", bold=True, color=WARN_TEXT if warn else EMERALD_TEXT)
    _run(para, text)
    _spacing(para)
    _spacing(doc.add_paragraph(), after=2)
    return table


def table(doc, header, rows, *, numeric=()):
    """numeric names the column indexes that hold figures; they align right."""
    grid = doc.add_table(rows=1, cols=len(header))
    _borders(grid, {"insideH": (LINE, 4), "bottom": (LINE, 4)})
    for i, name in enumerate(header):
        cell = grid.rows[0].cells[i]
        _shade(cell, EMERALD_BG)
        para = cell.paragraphs[0]
        _run(para, name, size=9.5, color=EMERALD_TEXT, bold=True)
        if i in numeric:
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in rows:
        cells = grid.add_row().cells
        for i, value in enumerate(row):
            para = cells[i].paragraphs[0]
            _run(para, str(value), size=9.5)
            if i in numeric:
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in grid.rows:
        for cell in row.cells:
            _cell_margins(cell, top=Mm(1.6), bottom=Mm(1.6), left=Mm(2.5), right=Mm(2.5))
            _spacing(cell.paragraphs[0])
    _spacing(doc.add_paragraph(), after=4)
    return grid


def _sample(path):
    doc = new_document("Document title", kind="Document type", meta="Author name · 23 October 2026")
    lead(doc, "One or two sentences that say what this document decides or reports.")
    heading(doc, "First section")
    paragraph(doc, "Body text in Geist. Keep paragraphs short and lead with the fact.")
    bullets(doc, ["A point the reader needs.", "A second point the reader needs."])
    callout(doc, "Key point.", "The one thing to remember from this section.")
    heading(doc, "Second section")
    table(doc, ["Item", "Owner", "Amount"], [["First row", "Team", "1,200"], ["Second row", "Team", "980"]],
          numeric=(2,))
    heading(doc, "A sub-section", level=3)
    paragraph(doc, "Sub-sections use level 3 and sit tight under their section.")
    callout(doc, "Open question.", "Something that still needs a decision.", warn=True)
    doc.add_page_break()
    heading(doc, "Third section")
    paragraph(doc, "Later pages carry the thin band at the top.")
    doc.save(path)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: margince_docx.py OUT.docx")
    _sample(sys.argv[1])
